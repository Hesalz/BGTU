
from pathlib import Path
from html import escape
import re
import os

from bs4 import BeautifulSoup
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUTPUT_DIR = Path("practical4_results")
OUTPUT_DIR.mkdir(exist_ok=True)

ZW0 = "\u200b"  # U+200B -> 0
ZW1 = "\u200c"  # U+200C -> 1
START = "1111111111111110"
END = "0111111111111111"

RUSSIAN_TEXT = (
    "Текстовая информация используется в электронных документах, сообщениях "
    "и веб-страницах. При передаче данных важно сохранять исходный вид "
    "документа и его читаемость. Стеганографические методы позволяют скрывать "
    "дополнительные сведения внутри обычного текста. В данной работе "
    "сравниваются два способа скрытого встраивания информации."
)
ENGLISH_TEXT = (
    "Textual information is widely used in electronic documents, messages, "
    "and web pages. During data transmission it is important to preserve the "
    "original appearance of a document and its readability. Steganographic "
    "methods make it possible to hide additional information inside ordinary "
    "text. This work compares two methods of hidden data embedding."
)

# Для практической работы носитель делаем заведомо длинным.
RUS_CARRIER = " ".join([RUSSIAN_TEXT] * 50)
ENG_CARRIER = " ".join([ENGLISH_TEXT] * 28)

MESSAGES = {
    "Русский": {
        "Короткое": "Секрет",
        "Среднее": "Это секретное сообщение для практического задания.",
        "Длинное": (
            "Это длинное секретное сообщение для проверки метода при "
            "увеличении объема скрываемых данных."
        ),
    },
    "English": {
        "Короткое": "Secret",
        "Среднее": "This is a secret message for the practical assignment.",
        "Длинное": (
            "This is a long secret message for testing the method with "
            "a larger amount of hidden data."
        ),
    },
}

def text_to_bits(text):
    return "".join(f"{byte:08b}" for byte in text.encode("utf-8"))

def bits_to_text(bits):
    if len(bits) % 8 != 0:
        return None
    try:
        return bytes(
            int(bits[i:i+8], 2) for i in range(0, len(bits), 8)
        ).decode("utf-8")
    except Exception:
        return None

def framed_bits(secret):
    return START + text_to_bits(secret) + END

# ============================================================
# SPACE STEGANOGRAPHY
# ============================================================

def space_encode(text, secret):
    bits = framed_bits(secret)
    words = text.split(" ")

    if len(words) - 1 < len(bits):
        raise ValueError("Недостаточно промежутков для Space Steganography.")

    result = []
    for i, word in enumerate(words[:-1]):
        result.append(word)
        result.append("  " if i < len(bits) and bits[i] == "1" else " ")
    result.append(words[-1])
    return "".join(result)

def space_decode(text):
    bits = ""
    for match in re.finditer(r" +", text):
        bits += "1" if len(match.group()) >= 2 else "0"

    start = bits.find(START)
    if start == -1:
        return None
    start += len(START)

    end = bits.find(END, start)
    if end == -1:
        return None

    return bits_to_text(bits[start:end])

def show_spaces(text, limit=700):
    return text[:limit].replace(" ", "·")

# ============================================================
# ZERO-WIDTH
# ============================================================

def zero_width_encode(text, secret):
    bits = framed_bits(secret)
    if len(text) < len(bits):
        raise ValueError("Недостаточно позиций для Zero-Width.")

    result = []
    for i, char in enumerate(text):
        result.append(char)
        if i < len(bits):
            result.append(ZW0 if bits[i] == "0" else ZW1)
    return "".join(result)

def zero_width_decode(text):
    bits = "".join(
        "0" if c == ZW0 else "1"
        for c in text if c in (ZW0, ZW1)
    )

    start = bits.find(START)
    if start == -1:
        return None
    start += len(START)

    end = bits.find(END, start)
    if end == -1:
        return None

    return bits_to_text(bits[start:end])

def show_zero_width(text, limit=700):
    return (
        text[:limit]
        .replace(ZW0, "[ZW0]")
        .replace(ZW1, "[ZW1]")
    )

# ============================================================
# HTML
# ============================================================

def save_html(text, path, title):
    html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>{escape(title)}</title>
</head>
<body>
<p>{escape(text)}</p>
</body>
</html>
"""
    path.write_text(html, encoding="utf-8")

def read_html(path):
    return BeautifulSoup(
        path.read_text(encoding="utf-8"), "html.parser"
    ).get_text()

# ============================================================
# PDF
# ============================================================

def find_font():
    candidates = []

    if os.name == "nt":
        candidates.extend([
            r"C:\Windows\Fonts\DejaVuSans.ttf",
            r"C:\Windows\Fonts\arial.ttf",
            r"C:\Windows\Fonts\segoeui.ttf",
            r"C:\Windows\Fonts\tahoma.ttf",
        ])
    else:
        candidates.extend([
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        ])

    for path in candidates:
        if Path(path).exists():
            return path

    return None

def prepare_pdf_font():
    font_path = find_font()
    if font_path:
        try:
            pdfmetrics.registerFont(TTFont("WorkFont", font_path))
            return "WorkFont"
        except Exception:
            pass
    return "Courier"

PDF_FONT = prepare_pdf_font()

def save_pdf(text, path):
    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    c.setFont(PDF_FONT, 7)
    y = height - 35

    # Не используем split() — сохраняем два пробела в записываемом тексте.
    max_chars = 105

    while text:
        line = text[:max_chars]
        text = text[max_chars:]
        c.drawString(25, y, line)
        y -= 10

        if y < 25:
            c.showPage()
            c.setFont(PDF_FONT, 7)
            y = height - 35

    c.save()

def read_pdf(path):
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)

# ============================================================
# ЭКСПЕРИМЕНТ
# ============================================================

def run_experiment(language, size, carrier, secret):
    space = space_encode(carrier, secret)
    zero = zero_width_encode(carrier, secret)

    space_html = OUTPUT_DIR / f"{language}_Space_{size}.html"
    zero_html = OUTPUT_DIR / f"{language}_ZW_{size}.html"
    space_pdf = OUTPUT_DIR / f"{language}_Space_{size}.pdf"
    zero_pdf = OUTPUT_DIR / f"{language}_ZW_{size}.pdf"

    save_html(space, space_html, f"{language} Space {size}")
    save_html(zero, zero_html, f"{language} Zero-Width {size}")
    save_pdf(space, space_pdf)
    save_pdf(zero, zero_pdf)

    html_space = read_html(space_html)
    html_zero = read_html(zero_html)
    pdf_space = read_pdf(space_pdf)
    pdf_zero = read_pdf(zero_pdf)

    result = {
        "language": language,
        "size": size,
        "message_chars": len(secret),
        "message_bytes": len(secret.encode("utf-8")),
        "message_bits": len(text_to_bits(secret)),
        "space_added_chars": len(space) - len(carrier),
        "zero_added_chars": len(zero) - len(carrier),
        "space_html": space_decode(html_space) == secret,
        "space_pdf": space_decode(pdf_space) == secret,
        "zero_html": zero_width_decode(html_zero) == secret,
        "zero_pdf": zero_width_decode(pdf_zero) == secret,
    }

    # Наглядность — только то, что нужно пользователю:
    print("\n" + "=" * 100)
    print(language, "|", size, "| Скрытое сообщение:", secret)
    print("=" * 100)

    print("\nИСХОДНЫЙ ТЕКСТ:")
    print(carrier[:700])

    print("\nSPACE — ПОСЛЕ ВСТРАИВАНИЯ")
    print("· = один пробел, ·· = два пробела")
    print(show_spaces(space))

    print("\nZERO-WIDTH — ПОСЛЕ ВСТРАИВАНИЯ")
    print("Обычный вид:")
    print(zero[:700])
    print("Для наглядности [ZW0] и [ZW1]:")
    print(show_zero_width(zero))

    return result

# ============================================================
# ГРАФИКИ
# ============================================================

def make_charts(results):
    # 1. Формат + метод.
    methods = ["Space HTML", "Space PDF", "Zero-Width HTML", "Zero-Width PDF"]
    values = [
        sum(r["space_html"] for r in results),
        sum(r["space_pdf"] for r in results),
        sum(r["zero_html"] for r in results),
        sum(r["zero_pdf"] for r in results),
    ]

    plt.figure(figsize=(9, 5))
    plt.bar(methods, values)
    plt.ylabel("Количество успешных извлечений")
    plt.title("Успешность извлечения скрытого сообщения")
    plt.ylim(0, 12)
    plt.tight_layout()
    plt.savefig(OUTPUT_DIR / "chart_success.png", dpi=160)
    plt.close()

    # 2. Увеличение размера текста.
    sizes = ["Короткое", "Среднее", "Длинное"]
    space_growth = []
    zero_growth = []

    for size in sizes:
        row = next(r for r in results if r["language"] == "Русский" and r["size"] == size)
        space_growth.append(row["space_added_chars"])
        zero_growth.append(row["zero_added_chars"])

    x = range(len(sizes))
    width = 0.36

    plt.figure(figsize=(9, 5))
    plt.bar([i - width/2 for i in x], space_growth, width=width, label="Space")
    plt.bar([i + width/2 for i in x], zero_growth, width=width, label="Zero-Width")
    plt.xticks(list(x), sizes)
    plt.ylabel("Добавлено символов")
    plt.title("Изменение размера текста после встраивания")
    plt.legend()
    plt.tight_layout()
    plt.savefig(OUTPUT_DIR / "chart_size.png", dpi=160)
    plt.close()

    # 3. Зависимость от длины сообщения — процент успешно восстановленных
    # экспериментов в каждой категории размера.
    space_by_size = []
    zero_by_size = []

    for size in sizes:
        rows = [r for r in results if r["size"] == size]
        space_by_size.append(
            100 * sum(r["space_html"] and r["space_pdf"] for r in rows) / len(rows)
        )
        zero_by_size.append(
            100 * sum(r["zero_html"] and r["zero_pdf"] for r in rows) / len(rows)
        )

    plt.figure(figsize=(9, 5))
    plt.plot(sizes, space_by_size, marker="o", label="Space")
    plt.plot(sizes, zero_by_size, marker="o", label="Zero-Width")
    plt.ylabel("Полное восстановление, %")
    plt.title("Зависимость результата от длины сообщения")
    plt.ylim(0, 100)
    plt.legend()
    plt.tight_layout()
    plt.savefig(OUTPUT_DIR / "chart_length.png", dpi=160)
    plt.close()

# ============================================================
# WORD-ОТЧЕТ
# ============================================================

def make_report(results):
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Практическое задание № 4")
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Сравнительный анализ методов текстовой стеганографии")
    r.bold = True

    doc.add_heading("1. Цель работы", level=1)
    doc.add_paragraph(
        "Провести сравнительный анализ метода замены пробелов "
        "(Space Steganography) и метода использования невидимых "
        "символов (Zero-Width Characters Steganography)."
    )

    doc.add_heading("2. Объекты и условия эксперимента", level=1)
    doc.add_paragraph(
        "Проверены русский и английский языки, форматы HTML и PDF, "
        "а также короткие, средние и длинные скрытые сообщения."
    )
    doc.add_paragraph(
        "В Space один пробел кодирует бит 0, два пробела — бит 1. "
        "В Zero-Width используются U+200B для 0 и U+200C для 1."
    )

    doc.add_heading("3. Результаты эксперимента", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        "Язык", "Размер", "Space HTML", "Space PDF",
        "ZW HTML", "ZW PDF"
    ]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    for r in results:
        row = table.add_row().cells
        vals = [
            r["language"],
            r["size"],
            "Да" if r["space_html"] else "Нет",
            "Да" if r["space_pdf"] else "Нет",
            "Да" if r["zero_html"] else "Нет",
            "Да" if r["zero_pdf"] else "Нет",
        ]
        for i, v in enumerate(vals):
            row[i].text = v

    doc.add_heading("4. Графическое представление результатов", level=1)
    for fname, caption in [
        ("chart_success.png", "Рисунок 1 — Успешность извлечения скрытого сообщения."),
        ("chart_size.png", "Рисунок 2 — Изменение размера текста при встраивании."),
        ("chart_length.png", "Рисунок 3 — Зависимость результата от длины сообщения."),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(OUTPUT_DIR / fname), width=Inches(6.0))
        p = doc.add_paragraph(caption)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("5. Сравнительный анализ", level=1)
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = "Table Grid"
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(["Критерий", "Space Steganography", "Zero-Width Characters"]):
        table2.rows[0].cells[i].text = h

    rows = [
        (
            "Степень заметности",
            "Изменение количества пробелов можно обнаружить при анализе текста.",
            "Невидимые символы не отображаются при обычном просмотре."
        ),
        (
            "Качество исходного текста",
            "Внешний вид почти не меняется, но количество пробелов изменяется.",
            "Внешний вид текста практически не меняется."
        ),
        (
            "Безопасность",
            "Зависит от сохранения структуры пробелов и возможности их анализа.",
            "Зависит от сохранения Unicode-символов и возможности их обнаружения."
        ),
        (
            "Практическая вместимость",
            "Определяется количеством промежутков между словами.",
            "Определяется количеством доступных позиций для вставки символов."
        ),
    ]

    for data in rows:
        cells = table2.add_row().cells
        for i, value in enumerate(data):
            cells[i].text = value

    doc.add_heading("6. Вывод", level=1)
    doc.add_paragraph(
        "Оба метода позволяют скрывать текстовую информацию. Space Steganography "
        "основан на изменении количества пробелов и поэтому требует сохранения "
        "исходной последовательности пробелов. Zero-Width использует невидимые "
        "Unicode-символы и практически не изменяет обычный вид текста."
    )
    doc.add_paragraph(
        "Результаты экспериментов показывают зависимость методов от формата "
        "документа и длины скрываемого сообщения. Выбор конкретного метода "
        "определяется требованиями к заметности, сохранению текста и обработке "
        "документа."
    )

    doc.save(OUTPUT_DIR / "Практическое_задание_4.docx")


def main():
    results = []

    for language, carrier in [
        ("Русский", RUS_CARRIER),
        ("English", ENG_CARRIER)
    ]:
        for size in ["Короткое", "Среднее", "Длинное"]:
            results.append(
                run_experiment(language, size, carrier, MESSAGES[language][size])
            )

    make_charts(results)
    make_report(results)

    print("\nИТОГОВАЯ ТАБЛИЦА")
    for r in results:
        print(
            r["language"], "|", r["size"],
            "| Space HTML:", r["space_html"],
            "| Space PDF:", r["space_pdf"],
            "| ZW HTML:", r["zero_html"],
            "| ZW PDF:", r["zero_pdf"]
        )

if __name__ == "__main__":
    main()
