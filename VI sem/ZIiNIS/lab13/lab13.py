import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


class TextSteganography:
    
    def __init__(self):
        self.END_MARKER = "####END####"
        
        self.KERNING_PAIRS = [
            'АВ', 'АГ', 'АЛ', 'АП', 'АТ', 'АФ', 'АЧ', 'АШ',
            'ВА', 'ВГ', 'ВД', 'ВЛ', 'ВП', 'ВТ', 'ВФ', 'ВХ', 'ВЦ', 'ВЧ', 'ВШ', 'ВЫ',
            'ГА', 'ГО', 'ГР', 'ГУ',
            'ДА', 'ДЕ', 'ДИ', 'ДЛ', 'ДО', 'ДУ',
            'ЖА', 'ЖЕ', 'ЖИ',
            'ЗА', 'ЗО',
            'КА', 'КЕ', 'КИ', 'КЛ', 'КО', 'КР', 'КУ',
            'ЛА', 'ЛЕ', 'ЛИ', 'ЛО', 'ЛУ',
            'МА', 'МЕ', 'МИ', 'МО', 'МУ',
            'НА', 'НЕ', 'НИ', 'НО', 'НУ',
            'ОА', 'ОВ', 'ОГ', 'ОД', 'ОК', 'ОЛ', 'ОМ', 'ОН', 'ОП', 'ОР', 'ОС', 'ОТ', 'ОФ', 'ОХ', 'ОЦ', 'ОЧ', 'ОШ', 'ОЩ',
            'ПА', 'ПЕ', 'ПИ', 'ПЛ', 'ПО', 'ПР', 'ПУ',
            'РА', 'РЕ', 'РИ', 'РО', 'РУ',
            'СА', 'СВ', 'СЕ', 'СИ', 'СК', 'СЛ', 'СО', 'СП', 'СР', 'СТ', 'СУ',
            'ТА', 'ТЕ', 'ТИ', 'ТО', 'ТР', 'ТУ',
            'УА', 'УВ', 'УГ', 'УД', 'УЖ', 'УЗ', 'УК', 'УЛ', 'УМ', 'УН', 'УП', 'УР', 'УС', 'УТ', 'УФ', 'УХ', 'УЦ', 'УЧ', 'УШ', 'УЩ',
            'ФА', 'ФЕ', 'ХА', 'ХЕ', 'ХИ', 'ХО', 'ЦА', 'ЦЕ', 'ЦИ',
            'ЧА', 'ЧЕ', 'ЧИ', 'ЧО', 'ША', 'ШЕ', 'ШИ', 'ШЛ', 'ШО',
            'ЩА', 'ЩЕ', 'ЫВ', 'ЫЙ', 'ЫЛ', 'ЫМ', 'ЫН', 'ЫП', 'ЫР', 'ЫС', 'ЫТ', 'ЫХ', 'ЫЧ',
            'ЬЕ', 'ЬИ', 'ЬО', 'ЬЯ', 'ЭТ', 'ЭФ',
            'ЮВ', 'ЮГ', 'ЮД', 'ЮЗ', 'ЮК', 'ЮЛ', 'ЮМ', 'ЮН', 'ЮП', 'ЮР', 'ЮС', 'ЮТ', 'ЮФ', 'ЮХ', 'ЮЦ', 'ЮЧ', 'ЮШ',
            'ЯВ', 'ЯГ', 'ЯД', 'ЯЗ', 'ЯК', 'ЯЛ', 'ЯМ', 'ЯН', 'ЯП', 'ЯР', 'ЯС', 'ЯТ', 'ЯФ', 'ЯХ', 'ЯЦ', 'ЯЧ', 'ЯШ', 'ЯЩ'
        ]
        
    def text_to_bits(self, text):
        text_bytes = text.encode('utf-8')
        bits = []
        for byte in text_bytes:
            for i in range(7, -1, -1):
                bits.append((byte >> i) & 1)
        return bits
    
    def bits_to_text(self, bits):
        bytes_list = []
        for i in range(0, len(bits), 8):
            if i + 8 <= len(bits):
                byte = 0
                for j in range(8):
                    byte = (byte << 1) | bits[i + j]
                bytes_list.append(byte)
        return bytes(bytes_list).decode('utf-8', errors='ignore')
    
    def get_capacity_spaces(self, text):
        """Емкость контейнера для пробелов"""
        words = re.findall(r'\S+', text)
        return len(words) - 1
    
    def embed_spaces(self, text, message, output_path):
        words = re.findall(r'\S+', text)
        full_message = message + self.END_MARKER
        bits = self.text_to_bits(full_message)
        
        available = len(words) - 1
        needed = len(bits)
        
        if needed > available:
            raise Exception(f"Не хватает пробелов! Нужно {needed}, доступно {available}")
        
        separators = re.findall(r'(\s+)', text)
        
        result_parts = []
        bit_index = 0
        
        for i, word in enumerate(words):
            result_parts.append(word)
            if i < len(words) - 1:
                if bit_index < len(bits):
                    if bits[bit_index] == 0:
                        result_parts.append(" ")
                    else:
                        result_parts.append("  ")
                    bit_index += 1
                else:
                    if i < len(separators):
                        result_parts.append(separators[i])
                    else:
                        result_parts.append(" ")
        
        result_text = ''.join(result_parts)
        
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(result_text)
        
        return bit_index // 8
    
    def extract_spaces(self, file_path):
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        
        words = re.findall(r'\S+', text)
        separators = re.findall(r'(\s+)', text)
        
        bits = []
        for i, sep in enumerate(separators[:len(words)-1]):
            space_count = len(re.findall(r' ', sep))
            if space_count == 1:
                bits.append(0)
            elif space_count == 2:
                bits.append(1)
        
        full_text = self.bits_to_text(bits)
        end_pos = full_text.find(self.END_MARKER)
        
        if end_pos != -1:
            return full_text[:end_pos]
        return "⚠️ Сообщение не найдено!"
    
    def find_kerning_pairs(self, text):
        """Находит все кернинговые пары в тексте"""
        words = text.split()
        pairs = []
        for word_idx, word in enumerate(words):
            for i in range(len(word) - 1):
                pair = word[i:i+2].upper()
                if pair in self.KERNING_PAIRS:
                    pairs.append({
                        'word_idx': word_idx,
                        'char_pos': i,
                        'word': word,
                        'pair': pair
                    })
        return pairs
    
    def get_capacity_kerning(self, text):
        """Емкость контейнера для кернинга"""
        return len(self.find_kerning_pairs(text))
    
    def embed_kerning(self, text, message, output_path):
        """Встраивание через кернинг (создает DOCX файл)"""
        pairs = self.find_kerning_pairs(text)
        full_message = message + self.END_MARKER
        bits = self.text_to_bits(full_message)
        
        available = len(pairs)
        needed = len(bits)
        
        if needed > available:
            raise Exception(f"Не хватает кернинговых пар! Нужно {needed}, найдено {available}")
        
        words = text.split()
        
        modifications = {}
        for i, pair_info in enumerate(pairs[:needed]):
            modifications[(pair_info['word_idx'], pair_info['char_pos'])] = bits[i]
        
        doc = Document()
        paragraph = doc.add_paragraph()
        
        for word_idx, word in enumerate(words):
            run = paragraph.add_run()
            for char_idx, char in enumerate(word):
                run.add_text(char)
                if (word_idx, char_idx) in modifications:
                    bit = modifications[(word_idx, char_idx)]
                    rPr = run._element.rPr
                    if rPr is None:
                        rPr = OxmlElement('w:rPr')
                        run._element.append(rPr)
                    spacing = OxmlElement('w:spacing')
                    if bit == 0:
                        spacing.set(qn('w:val'), '-3')
                    else:
                        spacing.set(qn('w:val'), '3')
                    rPr.append(spacing)
            
            if word_idx < len(words) - 1:
                run = paragraph.add_run()
                run.add_text(" ")
        
        doc.save(output_path)
        return len(bits) // 8
    
    def extract_kerning(self, file_path):
        bits = []
        
        with zipfile.ZipFile(file_path, 'r') as zip_ref:
            with zip_ref.open('word/document.xml') as xml_file:
                tree = ET.parse(xml_file)
                root = tree.getroot()
                
                namespaces = {
                    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
                }
                
                for spacing in root.findall('.//w:spacing', namespaces):
                    val = spacing.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')
                    if val:
                        try:
                            val_int = int(val)
                            if val_int < 0:
                                bits.append(0)
                            elif val_int > 0:
                                bits.append(1)
                        except:
                            pass
        
        full_text = self.bits_to_text(bits)
        end_pos = full_text.find(self.END_MARKER)
        
        if end_pos != -1:
            return full_text[:end_pos]
        return "⚠️ Сообщение не найдено! Проверьте, что вы выбрали правильный метод."


class SteganoGUI:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Текстовая стеганография - Вариант 3 (Пробелы + Кернинг)")
        self.root.geometry("1200x800")
        self.root.configure(bg='#2c3e50')
        
        self.stego = TextSteganography()
        self.setup_ui()
        self.load_default_texts()
        
    def setup_ui(self):
        main = tk.Frame(self.root, bg='#ecf0f1')
        main.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        left = tk.Frame(main, bg='#ecf0f1', width=480)
        left.pack(side=tk.LEFT, fill=tk.BOTH, padx=5, pady=5)
        left.pack_propagate(False)
        
        right = tk.Frame(main, bg='#ecf0f1')
        right.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        
        tk.Label(left, text="ТЕКСТОВАЯ СТЕГАНОГРАФИЯ", font=('Arial', 14, 'bold'),
                bg='#2c3e50', fg='white', pady=10).pack(fill=tk.X)
        
        frame_method = tk.LabelFrame(left, text="🔧 Выбор метода", bg='#ecf0f1', font=('Arial', 10, 'bold'))
        frame_method.pack(fill=tk.X, pady=10, padx=10)
        
        self.method_var = tk.StringVar(value="spaces")
        tk.Radiobutton(frame_method, text="Метод 1: Модификация числа пробелов (сохраняется в .txt)",
                      variable=self.method_var, value="spaces", bg='#ecf0f1').pack(anchor=tk.W, padx=20, pady=5)
        tk.Radiobutton(frame_method, text="Метод 2: Модификация кернинга (сохраняется в .docx)",
                      variable=self.method_var, value="kerning", bg='#ecf0f1').pack(anchor=tk.W, padx=20, pady=5)
        
        frame_container = tk.LabelFrame(left, text="📄 Текст-контейнер", 
                                        bg='#ecf0f1', font=('Arial', 10, 'bold'))
        frame_container.pack(fill=tk.BOTH, expand=True, pady=10, padx=10)
        
        btn_frame = tk.Frame(frame_container, bg='#ecf0f1')
        btn_frame.pack(fill=tk.X, pady=5)
        tk.Button(btn_frame, text="Загрузить .txt", command=self.load_file,
                 bg='#3498db', fg='white').pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="Стандартный текст", command=self.load_default_container,
                 bg='#f39c12', fg='white').pack(side=tk.LEFT, padx=5)
        
        self.capacity_label = tk.Label(frame_container, text="", bg='#ecf0f1', font=('Arial', 9))
        self.capacity_label.pack(pady=5)
        
        self.container_text = scrolledtext.ScrolledText(frame_container, height=12, font=('Arial', 10))
        self.container_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        self.container_text.bind('<KeyRelease>', self.update_capacity)
        
        frame_msg = tk.LabelFrame(left, text="💬 Секретное сообщение", 
                                   bg='#ecf0f1', font=('Arial', 10, 'bold'))
        frame_msg.pack(fill=tk.BOTH, expand=True, pady=10, padx=10)
        
        notebook = ttk.Notebook(frame_msg)
        notebook.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        self.fio_text = scrolledtext.ScrolledText(notebook, height=4, font=('Arial', 10))
        self.report_text = scrolledtext.ScrolledText(notebook, height=4, font=('Arial', 10))
        self.custom_text = scrolledtext.ScrolledText(notebook, height=4, font=('Arial', 10))
        
        notebook.add(self.fio_text, text="📝 ФИО")
        notebook.add(self.report_text, text="📄 Текст отчета")
        notebook.add(self.custom_text, text="✏️ Свой текст")
        
        btn_actions = tk.Frame(left, bg='#ecf0f1')
        btn_actions.pack(fill=tk.X, pady=10, padx=10)
        
        tk.Button(btn_actions, text="▶ ВСТРОИТЬ", command=self.embed,
                 bg='#27ae60', fg='white', font=('Arial', 12, 'bold'), pady=8).pack(fill=tk.X, pady=5)
        
        tk.Button(btn_actions, text="◀ ИЗВЛЕЧЬ", command=self.extract,
                 bg='#e67e22', fg='white', font=('Arial', 12, 'bold'), pady=8).pack(fill=tk.X, pady=5)
        
        frame_result = tk.LabelFrame(right, text="📋 Результат", bg='#ecf0f1', font=('Arial', 10, 'bold'))
        frame_result.pack(fill=tk.BOTH, expand=True, pady=5)
        
        self.result_text = scrolledtext.ScrolledText(frame_result, height=20, bg='#e8f8f5', font=('Arial', 11))
        self.result_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        frame_info = tk.LabelFrame(right, text="ℹ️ Информация", bg='#ecf0f1', font=('Arial', 10, 'bold'))
        frame_info.pack(fill=tk.X, pady=5)
        
        self.info_text = scrolledtext.ScrolledText(frame_info, height=6, bg='#fef9e7')
        self.info_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        self.status_var = tk.StringVar(value="✅ Готов к работе")
        status = tk.Label(self.root, textvariable=self.status_var, bg='#2c3e50', fg='white',
                         font=('Arial', 9), pady=5)
        status.pack(side=tk.BOTTOM, fill=tk.X)
        
        self.update_capacity()
        self.update_info()
    
    def update_info(self):
        info = """【 МЕТОД 1: ПРОБЕЛЫ 】
Правило: 1 пробел = бит 0, 2 пробела = бит 1
Сохраняется в .txt файл

【 МЕТОД 2: КЕРНИНГ 】
Правило: отрицательный интервал = бит 0, положительный = бит 1
Сохраняется в .docx файл (Microsoft Word)

【 КЕРНИНГОВЫЕ ПАРЫ 】
Сочетания букв, которые визуально выглядят с расстоянием:
АВ, АГ, АЛ, АП, АТ, АФ, ВА, ВГ, ВД, ВЛ, ВП, ВТ, ГА, ДА, ЖА, КА, ЛА, МА, НА, ОВ, ПА, РА, СА, ТА, УА, ФА, ХА, ЦА, ЧА, ША, ЩА, ЯВ и др."""
        
        self.info_text.config(state=tk.NORMAL)
        self.info_text.delete("1.0", tk.END)
        self.info_text.insert("1.0", info)
        self.info_text.config(state=tk.DISABLED)
    
    def update_capacity(self, event=None):
        text = self.container_text.get("1.0", tk.END).strip()
        message = self.get_current_message()
        
        if text:
            spaces_capacity = self.stego.get_capacity_spaces(text)
            kerning_capacity = self.stego.get_capacity_kerning(text)
        else:
            spaces_capacity = 0
            kerning_capacity = 0
        
        if message:
            full_message = message + self.stego.END_MARKER
            needed_bits = len(self.stego.text_to_bits(full_message))
        else:
            needed_bits = 0
        
        words_count = len(re.findall(r'\S+', text)) if text else 0
        
        status = f"📊 Слов: {words_count} | Пробелов: {spaces_capacity} | Кернинг пар: {kerning_capacity} | Нужно бит: {needed_bits}"
        
        if needed_bits <= spaces_capacity:
            status += " | Пробелы: OK"
        else:
            status += f" | Пробелы: не хватает {needed_bits - spaces_capacity}"
        
        if needed_bits <= kerning_capacity:
            status += " | Кернинг: OK"
        else:
            status += f" | Кернинг: не хватает {needed_bits - kerning_capacity}"
        
        self.capacity_label.config(text=status, fg='green' if needed_bits <= max(spaces_capacity, kerning_capacity) else 'red')
    
    def get_current_message(self):
        return self.fio_text.get("1.0", tk.END).strip()
    
    def load_default_container(self):
        """Загружает большой текст-контейнер"""
        text = """Информационная безопасность является одной из важнейших проблем современного общества Защита данных от несанкционированного доступа требует применения различных методов криптографии и стеганографии В данной лабораторной работе исследуются методы текстовой стеганографии на основе модификации пробелов и кернинга Эти методы позволяют скрывать информацию в текстовых документах без изменения визуального восприятия текста

Криптография занимается шифрованием информации делая её недоступной для чтения без ключа Стеганография же скрывает сам факт передачи секретного сообщения маскируя его под обычный текст В современном мире оба подхода часто используются вместе для максимальной защиты данных

Метод модификации пробелов является одним из простейших методов текстовой стеганографии Он заключается в изменении количества пробелов между словами текста контейнера Один пробел кодирует бит ноль два пробела кодируют бит один Человеческий глаз не замечает разницы между одним и двумя пробелами

Метод модификации кернинга более сложный Кернинг это изменение межбуквенного интервала для определенных пар символов В русском языке существует множество кернинговых пар АВ АГ АЛ АП АТ АФ АЧ АШ ВА ВГ ВД ВЛ ВП ВТ и многие другие Изменяя интервал между буквами в таких парах можно кодировать биты сообщения Положительный кернинг кодирует единицу отрицательный кернинг кодирует ноль

Важным требованием к тексту контейнеру является его достаточный объем Чем больше слов и кернинговых пар в тексте тем больше секретного сообщения можно в него встроить Разработанное приложение позволяет встраивать секретные сообщения в текстовые документы и извлекать их обратно Пользователь может выбрать один из двух методов стеганографии и сохранить результат в отдельный файл"""
        
        self.container_text.delete("1.0", tk.END)
        self.container_text.insert("1.0", text)
        self.update_status("Загружен стандартный текст")
        self.update_capacity()
    
    def load_file(self):
        path = filedialog.askopenfilename(filetypes=[("Text files", "*.txt")])
        if path:
            try:
                with open(path, 'r', encoding='utf-8') as f:
                    text = f.read()
                self.container_text.delete("1.0", tk.END)
                self.container_text.insert("1.0", text)
                self.update_status(f"Загружен: {os.path.basename(path)}")
                self.update_capacity()
            except Exception as e:
                messagebox.showerror("Ошибка", str(e))
    
    def load_default_texts(self):
        """Загружает примеры сообщений"""
        self.load_default_container()
        self.fio_text.insert("1.0", "Иванов Иван Иванович")
        self.report_text.insert("1.0", "ЛАБОРАТОРНАЯ РАБОТА №13\nИсследование методов текстовой стеганографии\n\nЦель: Изучение методов модификации пробелов и кернинга.\n\nВывод: Методы эффективны для сокрытия информации в текстовых документах.")
        self.update_capacity()
    
    def embed(self):
        """Встраивание сообщения"""
        text = self.container_text.get("1.0", tk.END).strip()
        if not text:
            messagebox.showerror("Ошибка", "Введите текст-контейнер!")
            return
        
        message = self.get_current_message()
        if not message:
            messagebox.showerror("Ошибка", "Введите сообщение!")
            return
        
        method = self.method_var.get()
        
        if method == "spaces":
            output_path = filedialog.asksaveasfilename(
                defaultextension=".txt",
                filetypes=[("Text files", "*.txt")]
            )
        else:
            output_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Word documents", "*.docx")]
            )
        
        if not output_path:
            return
        
        try:
            self.update_status("Встраивание...")
            
            if method == "spaces":
                bytes_embedded = self.stego.embed_spaces(text, message, output_path)
                method_name = "ПРОБЕЛЫ"
            else:
                bytes_embedded = self.stego.embed_kerning(text, message, output_path)
                method_name = "КЕРНИНГ"
            
            self.result_text.delete("1.0", tk.END)
            self.result_text.insert("1.0", 
                f"✅ ВСТРОЕНО УСПЕШНО!\n\n"
                f"Метод: {method_name}\n"
                f"Сообщение: '{message}'\n"
                f"Байт: {bytes_embedded}\n"
                f"Файл: {output_path}\n\n"
                f"Теперь нажмите 'ИЗВЛЕЧЬ' и выберите этот же файл.")
            
            self.update_status(f"✅ Встроено {bytes_embedded} байт методом {method_name}")
            messagebox.showinfo("Успех", f"Сообщение встроено!\nМетод: {method_name}\nФайл: {output_path}")
            
        except Exception as e:
            messagebox.showerror("Ошибка", str(e))
            self.update_status("❌ Ошибка встраивания")
    
    def extract(self):
        """Извлечение сообщения"""
        method = self.method_var.get()
        
        if method == "spaces":
            path = filedialog.askopenfilename(filetypes=[("Text files", "*.txt")])
            if not path:
                return
            
            try:
                self.update_status("Извлечение...")
                result = self.stego.extract_spaces(path)
                self.result_text.delete("1.0", tk.END)
                self.result_text.insert("1.0", f"📩 ИЗВЛЕЧЕННОЕ СООБЩЕНИЕ:\n\n{result}")
                self.update_status(f"✅ Извлечено {len(result)} символов")
            except Exception as e:
                messagebox.showerror("Ошибка", str(e))
                
        else: 
            path = filedialog.askopenfilename(filetypes=[("Word documents", "*.docx")])
            if not path:
                return
            
            try:
                self.update_status("Извлечение...")
                result = self.stego.extract_kerning(path)
                self.result_text.delete("1.0", tk.END)
                self.result_text.insert("1.0", f"📩 ИЗВЛЕЧЕННОЕ СООБЩЕНИЕ:\n\n{result}")
                self.update_status(f"✅ Извлечено {len(result)} символов")
            except Exception as e:
                messagebox.showerror("Ошибка", str(e))
    
    def update_status(self, msg):
        self.status_var.set(msg)
        print(f"[STATUS] {msg}")
    
    def run(self):
        self.root.mainloop()


if __name__ == "__main__":
    app = SteganoGUI()
    app.run()